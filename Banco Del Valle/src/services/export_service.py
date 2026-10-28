"""
Servicio de exportación documental multiformato (CSV, Excel .xlsx y PDF membretado).
Genera extractos bancarios, cronogramas franceses y reportes ejecutivos de tesorería.
Referencia de diseño: docs/diagrama_logica_negocio_banco.md
"""

from datetime import datetime
import csv
import io
import os
from typing import Any, Dict, List, Optional, Tuple
import fitz
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


def _obtener_ruta_logo_png() -> Optional[str]:
    """Obtiene o genera bajo demanda el isotipo oficial B+V en formato PNG."""
    base_dir = os.path.dirname(os.path.abspath(__file__))
    candidatos_png = [
        os.path.join(base_dir, "..", "..", "frontend", "public", "logo-valle.png"),
        os.path.join(base_dir, "..", "..", "frontend", "dist", "logo-valle.png"),
        os.path.abspath("Banco Del Valle/frontend/public/logo-valle.png"),
    ]
    for c in candidatos_png:
        if os.path.exists(c):
            return os.path.abspath(c)

    # Si no existe PNG pero sí el SVG oficial, renderizarlo a PNG 300 DPI con transparencia alfa
    candidatos_svg = [
        os.path.join(base_dir, "..", "..", "frontend", "public", "logo-valle.svg"),
        os.path.join(base_dir, "..", "..", "frontend", "dist", "logo-valle.svg"),
        os.path.abspath("Banco Del Valle/frontend/public/logo-valle.svg"),
    ]
    for s in candidatos_svg:
        if os.path.exists(s):
            try:
                doc_svg = fitz.open(s)
                pix = doc_svg[0].get_pixmap(alpha=True, dpi=300)
                png_salida = os.path.splitext(s)[0] + ".png"
                pix.save(png_salida)
                return os.path.abspath(png_salida)
            except Exception:
                pass
    return None


def exportar_a_csv(datos: List[Dict[str, Any]], campos: Optional[List[str]] = None) -> bytes:
    """Exporta una lista de diccionarios a formato CSV con codificación UTF-8-SIG para Excel."""
    if not datos:
        return b""

    keys = campos if campos else list(datos[0].keys())
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=keys, extrasaction="ignore", delimiter=";")
    writer.writeheader()
    for fila in datos:
        writer.writerow(fila)

    return output.getvalue().encode("utf-8-sig")


def exportar_a_excel(
    datos: List[Dict[str, Any]],
    nombre_hoja: str = "Reporte",
    titulo_reporte: str = "BANCO DEL VALLE — REPORTE INSTITUCIONAL",
    subtitulo: Optional[str] = None,
    kpis: Optional[Dict[str, Any]] = None,
) -> bytes:
    """Exporta registros a una planilla Excel (.xlsx) con estilos bancarios institucionales y logo."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = nombre_hoja[:30]

    # Paleta Catamarca: Azul Marino #0F3A5D, Naranja #FF7A00
    fill_header = PatternFill(start_color="0F3A5D", end_color="0F3A5D", fill_type="solid")
    fill_kpi = PatternFill(start_color="F5F7FA", end_color="F5F7FA", fill_type="solid")
    font_header = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    font_banco = Font(name="Segoe UI", size=15, bold=True, color="0F3A5D")
    font_sub = Font(name="Segoe UI", size=9, italic=True, color="555555")
    font_titulo = Font(name="Segoe UI", size=12, bold=True, color="FF7A00")
    font_kpi_label = Font(name="Segoe UI", size=8, bold=True, color="555555")
    font_kpi_val = Font(name="Segoe UI", size=10, bold=True, color="0F3A5D")
    font_datos = Font(name="Segoe UI", size=10)

    thin_border = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )

    # 1. Logo institucional en A1
    logo_path = _obtener_ruta_logo_png()
    if logo_path and os.path.exists(logo_path):
        try:
            from openpyxl.drawing.image import Image as ExcelImage
            img = ExcelImage(logo_path)
            img.width = 46
            img.height = 46
            ws.add_image(img, "A1")
        except Exception:
            pass

    # 2. Encabezado institucional en columna B
    ws["B1"] = "BANCO DEL VALLE"
    ws["B1"].font = font_banco
    ws["B2"] = "San Fernando del Valle de Catamarca — Entidad Regulada BCRA"
    ws["B2"].font = font_sub
    ws["B3"] = titulo_reporte
    ws["B3"].font = font_titulo

    emision_str = f"Fecha y Hora de Emisión: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} - Catamarca, Argentina"
    if subtitulo:
        ws["B4"] = f"{subtitulo} | {emision_str}"
    else:
        ws["B4"] = emision_str
    ws["B4"].font = font_sub

    ws.row_dimensions[1].height = 20
    ws.row_dimensions[2].height = 15
    ws.row_dimensions[3].height = 18
    ws.row_dimensions[4].height = 16

    current_row = 6

    # 3. Bloque de KPIs si están presentes
    if kpis:
        ws.row_dimensions[current_row].height = 16
        ws.row_dimensions[current_row + 1].height = 22
        for col_idx, (k, v) in enumerate(kpis.items(), start=1):
            c_lbl = ws.cell(row=current_row, column=col_idx, value=k.upper())
            c_lbl.font = font_kpi_label
            c_lbl.fill = fill_kpi
            c_lbl.alignment = Alignment(horizontal="center", vertical="center")
            c_lbl.border = thin_border

            c_val = ws.cell(row=current_row + 1, column=col_idx)
            c_val.font = font_kpi_val
            c_val.fill = fill_kpi
            c_val.alignment = Alignment(horizontal="center", vertical="center")
            c_val.border = thin_border

            if isinstance(v, float):
                c_val.value = v
                if "TNA" in k.upper() or "TASA" in k.upper():
                    c_val.number_format = "0.0%"
                else:
                    c_val.number_format = "$#,##0.00"
            elif isinstance(v, int):
                c_val.value = v
                c_val.number_format = "#,##0"
            else:
                c_val.value = str(v)

        current_row += 3
    else:
        current_row += 1

    if not datos:
        ws.cell(row=current_row, column=1, value="No se encontraron registros para el período solicitado.")
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    # 4. Cabecera de la tabla de datos
    headers = list(datos[0].keys())
    header_row_idx = current_row
    ws.row_dimensions[header_row_idx].height = 24

    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=header_row_idx, column=col_idx)
        cell.value = h.replace("_", " ").upper()
        cell.fill = fill_header
        cell.font = font_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # 5. Filas de datos
    for fila in datos:
        current_row += 1
        ws.row_dimensions[current_row].height = 19
        for col_idx, h in enumerate(headers, start=1):
            val = fila.get(h)
            c = ws.cell(row=current_row, column=col_idx, value=val if val is not None else "")
            c.font = font_datos
            c.border = thin_border

            # Formatos numéricos específicos
            h_upper = h.upper()
            if "NUMERO" in h_upper or h_upper == "CUOTA" or h_upper == "ID":
                c.alignment = Alignment(horizontal="center", vertical="center")
                if isinstance(val, (int, float)):
                    c.number_format = "0"
            elif isinstance(val, float):
                c.number_format = "$#,##0.00"
            elif isinstance(val, int):
                c.number_format = "#,##0"

    # 6. Ajuste generoso de ancho de columnas para garantizar legibilidad total
    for col in ws.columns:
        col_letter = col[0].column_letter
        max_len = 0
        for cell in col:
            # Medir únicamente las celdas de la tabla (desde header_row_idx)
            if cell.row < header_row_idx:
                continue
            val_text = str(cell.value or "")
            if len(val_text) > max_len:
                max_len = len(val_text)
        
        # Columna de cuota necesita ~18, las monetarias ~25 a 30
        if col_letter == "A":
            ws.column_dimensions[col_letter].width = max(max_len + 4, 18)
        else:
            ws.column_dimensions[col_letter].width = max(max_len + 6, 26)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def exportar_a_pdf(
    titulo: str,
    subtitulo: str,
    columnas: List[str],
    filas: List[List[Any]],
    kpis: Optional[Dict[str, Any]] = None,
) -> bytes:
    """Genera un documento PDF institucional con membrete, tabla formateada y paginación."""
    doc = fitz.open()
    ancho_pagina, alto_pagina = 595, 842  # A4 vertical
    color_primario = (15 / 255, 58 / 255, 93 / 255)      # #0F3A5D Azul Marino
    color_secundario = (255 / 255, 122 / 255, 0 / 255)  # #FF7A00 Naranja
    color_gris = (245 / 255, 247 / 255, 250 / 255)

    def _crear_nueva_pagina():
        page = doc.new_page(width=ancho_pagina, height=alto_pagina)
        # Membrete superior azul
        page.draw_rect(fitz.Rect(0, 0, ancho_pagina, 65), fill=color_primario, color=None)
        # Línea de realce naranja
        page.draw_rect(fitz.Rect(0, 65, ancho_pagina, 68), fill=color_secundario, color=None)

        # Inclusión del isotipo B+V oficial
        logo_path = _obtener_ruta_logo_png()
        if logo_path and os.path.exists(logo_path):
            try:
                page.insert_image(fitz.Rect(35, 12, 75, 52), filename=logo_path)
                page.insert_text((85, 36), "BANCO DEL VALLE", fontsize=17, fontname="helv", color=(1, 1, 1))
                page.insert_text(
                    (85, 50),
                    "San Fernando del Valle de Catamarca — Entidad Regulada BCRA",
                    fontsize=8,
                    fontname="helv",
                    color=(0.85, 0.9, 0.95),
                )
            except Exception:
                page.insert_text((35, 38), "BANCO DEL VALLE", fontsize=18, fontname="helv", color=(1, 1, 1))
                page.insert_text(
                    (35, 52),
                    "San Fernando del Valle de Catamarca — Entidad Regulada BCRA",
                    fontsize=8,
                    fontname="helv",
                    color=(0.85, 0.9, 0.95),
                )
        else:
            page.insert_text((35, 38), "BANCO DEL VALLE", fontsize=18, fontname="helv", color=(1, 1, 1))
            page.insert_text(
                (35, 52),
                "San Fernando del Valle de Catamarca — Entidad Regulada BCRA",
                fontsize=8,
                fontname="helv",
                color=(0.85, 0.9, 0.95),
            )
        return page

    pagina_actual = _crear_nueva_pagina()
    y_pos = 95

    # Título del reporte
    pagina_actual.insert_text((35, y_pos), titulo.upper(), fontsize=13, fontname="helv", color=color_primario)
    y_pos += 15
    pagina_actual.insert_text((35, y_pos), subtitulo, fontsize=9, fontname="helv", color=(0.3, 0.3, 0.3))
    y_pos += 20

    # Bloque de KPIs destacados si están presentes
    if kpis:
        # Ponderación proporcional de anchos según tipo de KPI (para evitar desbordes)
        pesos_kpi = {"Titular": 1.6, "Capital": 1.0, "TNA": 0.65, "Plazo": 0.75, "Cuota Fija": 1.0}
        total_peso = sum(pesos_kpi.get(k, 1.0) for k in kpis)
        ancho_disponible = ancho_pagina - 70

        x_actual = 35.0
        for clave, valor in kpis.items():
            peso = pesos_kpi.get(clave, 1.0)
            caja_w = (peso / total_peso) * ancho_disponible
            box_rect = fitz.Rect(x_actual, y_pos, x_actual + caja_w - 6, y_pos + 42)
            pagina_actual.draw_rect(box_rect, fill=color_gris, color=(0.8, 0.85, 0.9))
            pagina_actual.insert_text(
                (box_rect.x0 + 6, box_rect.y0 + 13), clave.upper(), fontsize=7, color=(0.4, 0.4, 0.4)
            )

            # Auto-ajuste de tamaño y truncado limpio para no desbordar la caja
            val_str = str(valor)
            max_ancho_texto = box_rect.width - 12
            fs = 10.0
            while fs > 7.5 and fitz.get_text_length(val_str, fontname="helv", fontsize=fs) > max_ancho_texto:
                fs -= 0.5
            if fitz.get_text_length(val_str, fontname="helv", fontsize=fs) > max_ancho_texto:
                while len(val_str) > 4 and fitz.get_text_length(val_str + "...", fontname="helv", fontsize=fs) > max_ancho_texto:
                    val_str = val_str[:-1]
                val_str = val_str.rstrip() + "..."

            pagina_actual.insert_text(
                (box_rect.x0 + 6, box_rect.y0 + 30), val_str, fontsize=fs, fontname="helv", color=color_primario
            )
            x_actual += caja_w
        y_pos += 55

    # Renderizado de la tabla de datos
    num_cols = len(columnas)
    col_width = (ancho_pagina - 70) / num_cols

    # Cabecera de la tabla
    header_rect = fitz.Rect(35, y_pos, ancho_pagina - 35, y_pos + 20)
    pagina_actual.draw_rect(header_rect, fill=color_primario, color=None)
    for i, col_name in enumerate(columnas):
        pagina_actual.insert_text(
            (35 + i * col_width + 4, y_pos + 14),
            col_name[:18].upper(),
            fontsize=8,
            fontname="helv",
            color=(1, 1, 1),
        )
    y_pos += 22

    # Filas de la tabla
    for r_idx, fila in enumerate(filas):
        if y_pos > alto_pagina - 50:
            pagina_actual = _crear_nueva_pagina()
            y_pos = 90

        fila_rect = fitz.Rect(35, y_pos, ancho_pagina - 35, y_pos + 17)
        if r_idx % 2 == 1:
            pagina_actual.draw_rect(fila_rect, fill=color_gris, color=None)

        for c_idx, valor in enumerate(fila):
            texto = f"{valor:,.2f}" if isinstance(valor, float) else str(valor)
            pagina_actual.insert_text(
                (35 + c_idx * col_width + 4, y_pos + 12),
                texto[:22],
                fontsize=7.5,
                fontname="helv",
                color=(0.15, 0.15, 0.15),
            )
        y_pos += 18

    # Pie de página en todas las hojas
    for num_p, p in enumerate(doc):
        p.insert_text(
            (35, alto_pagina - 20),
            f"Documento generado el {datetime.now().strftime('%d/%m/%Y %H:%M')} — Página {num_p + 1} de {len(doc)}",
            fontsize=7,
            color=(0.5, 0.5, 0.5),
        )

    return doc.tobytes()


def exportar_extracto_cuenta(
    movimientos: List[Dict[str, Any]],
    formato: str = "pdf",
    titular: str = "Juan José Cano",
    cbu: str = "2850001211100012345678",
    alias: str = "JUAN.BANCO.VALLE",
    saldo_actual: float = 3000000.0,
) -> Tuple[bytes, str, str]:
    """Genera el extracto bancario oficial en el formato solicitado (pdf, excel, csv)."""
    formato_limpio = formato.lower().strip()
    fecha_archivo = datetime.now().strftime("%Y%m%d_%H%M")

    if formato_limpio == "csv":
        archivo_bytes = exportar_a_csv(movimientos)
        return archivo_bytes, "text/csv", f"extracto_{alias}_{fecha_archivo}.csv"

    elif formato_limpio in ("excel", "xlsx"):
        archivo_bytes = exportar_a_excel(
            datos=movimientos,
            nombre_hoja="Extracto Bancario",
            titulo_reporte=f"EXTRACTO BANCARIO — TITULAR: {titular} (CBU: {cbu})",
        )
        return (
            archivo_bytes,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            f"extracto_{alias}_{fecha_archivo}.xlsx",
        )

    # Por defecto PDF
    columnas = ["FECHA", "OPERACIÓN", "DETALLE", "MONTO", "SALDO POSTERIOR"]
    filas = [
        [
            m.get("fecha_hora", "")[:16],
            m.get("tipo_movimiento", "").replace("_", " "),
            m.get("descripcion", ""),
            f"${float(m.get('monto', 0.0)):,.2f}",
            f"${float(m.get('saldo_posterior', 0.0)):,.2f}",
        ]
        for m in movimientos
    ]
    kpis = {
        "Titular": titular,
        "Alias": alias,
        "Saldo Actual": f"${saldo_actual:,.2f}",
        "Movimientos": len(movimientos),
    }
    archivo_bytes = exportar_a_pdf(
        titulo="Resumen de Cuenta Transaccional Remunerada",
        subtitulo=f"CBU: {cbu} | Sucursal San Fernando del Valle de Catamarca",
        columnas=columnas,
        filas=filas,
        kpis=kpis,
    )
    return archivo_bytes, "application/pdf", f"extracto_{alias}_{fecha_archivo}.pdf"


def exportar_cronograma_prestamo(
    cronograma: List[Dict[str, Any]],
    monto_capital: float,
    tna: float,
    cantidad_cuotas: int,
    cuota_fija: float,
    formato: str = "pdf",
    titular: str = "Juan José Cano",
) -> Tuple[bytes, str, str]:
    """Genera la cotización y tabla de amortización francesa de un crédito."""
    formato_limpio = formato.lower().strip()
    fecha_archivo = datetime.now().strftime("%Y%m%d")

    if formato_limpio == "csv":
        archivo_bytes = exportar_a_csv(cronograma)
        return archivo_bytes, "text/csv", f"cotizacion_credito_{fecha_archivo}.csv"

    elif formato_limpio in ("excel", "xlsx"):
        kpis_excel = {
            "Titular": titular,
            "Capital Solicitado": monto_capital,
            "TNA Pactada": tna,
            "Cantidad de Cuotas": f"{cantidad_cuotas} cuotas",
            "Cuota Mensual": cuota_fija,
        }
        archivo_bytes = exportar_a_excel(
            datos=cronograma,
            nombre_hoja="Cronograma Frances",
            titulo_reporte="COTIZACIÓN OFICIAL — PLAN DE AMORTIZACIÓN FRANCESA",
            subtitulo=f"Titular: {titular} — Capital: ${monto_capital:,.2f}",
            kpis=kpis_excel,
        )
        return (
            archivo_bytes,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            f"cotizacion_credito_{fecha_archivo}.xlsx",
        )

    columnas = ["CUOTA N°", "MONTO CUOTA", "CAPITAL", "INTERÉS", "SALDO RESTANTE"]
    filas = [
        [
            c.get("numero_cuota", 0),
            f"${float(c.get('monto_cuota', 0.0)):,.2f}",
            f"${float(c.get('capital_amortizado', 0.0)):,.2f}",
            f"${float(c.get('interes_cuota', 0.0)):,.2f}",
            f"${float(c.get('saldo_capital_restante', 0.0)):,.2f}",
        ]
        for c in cronograma
    ]
    kpis = {
        "Titular": titular,
        "Capital": f"${monto_capital:,.2f}",
        "TNA": f"{tna * 100:.1f}%",
        "Plazo": f"{cantidad_cuotas} meses",
        "Cuota Fija": f"${cuota_fija:,.2f}",
    }
    archivo_bytes = exportar_a_pdf(
        titulo="Cotización de Préstamo — Sistema Francés",
        subtitulo="Simulación oficial con amortización de capital creciente e interés decreciente",
        columnas=columnas,
        filas=filas,
        kpis=kpis,
    )
    return archivo_bytes, "application/pdf", f"cotizacion_credito_{fecha_archivo}.pdf"
