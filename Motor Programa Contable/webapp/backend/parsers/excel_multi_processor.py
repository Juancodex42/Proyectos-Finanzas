import io
import re
import datetime
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd

try:
    from .base_parser import (
        _normalize_text, parse_date, clean_decimal,
        create_transaction, normalize_exchange_name, get_usd_ars_rate
    )
except ImportError:
    from parsers.base_parser import (
        _normalize_text, parse_date, clean_decimal,
        create_transaction, normalize_exchange_name, get_usd_ars_rate
    )

_STABLES = {'USDT', 'USDC', 'DAI', 'BUSD', 'USD', 'USDP', 'FDUSD', 'PYUSD', 'UXD', 'TUSD'}
_REF_USD = {
    'BTC': 60000.0, 'ETH': 3000.0, 'SOL': 150.0, 'BNB': 600.0,
    'AVAX': 30.0, 'LINK': 15.0, 'DOT': 7.0, 'XRP': 0.6, 'ADA': 0.5,
    'MATIC': 0.5, 'POL': 0.5, 'LTC': 70.0, 'BCH': 350.0, 'DOGE': 0.1,
    'SHIB': 0.00002, 'TRX': 0.12, 'NEAR': 5.0, 'ATOM': 5.0, 'UNI': 7.0
}


def _get_usd_ars_rate_for_date(fecha_dt):
    return get_usd_ars_rate(fecha_dt)


def generate_master_excel(transactions):
    """Generates the master consolidated Excel workbook styled by exchange."""
    if not transactions:
        return None
    
    output = io.BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = "Consolidado"
    
    STYLES = {
        'Binance Spot': {'fill': 'F3BA2F', 'font': '000000'},
        'Binance P2P': {'fill': 'F3BA2F', 'font': '000000'},
        'Binance': {'fill': 'F3BA2F', 'font': '000000'},
        'Bitso': {'fill': '29AC00', 'font': 'FFFFFF'},
        'Fiwind': {'fill': 'EFB41D', 'font': '000000'},
        'Ripio Trade': {'fill': '7A3CE3', 'font': 'FFFFFF'},
        'Ripio Classic': {'fill': '7A3CE3', 'font': 'FFFFFF'},
        'Ripio': {'fill': '7A3CE3', 'font': 'FFFFFF'},
        'OKX': {'fill': '111111', 'font': 'FFFFFF'},
        'Bybit': {'fill': 'F7A600', 'font': '000000'},
        'Bitget': {'fill': '00F0FF', 'font': '191F61'},
        'Generic': {'fill': '4B5563', 'font': 'FFFFFF'}
    }
    
    TYPE_STYLES = {
        'Compra': {'color': '2E7D32'},
        'Venta': {'color': 'C62828'},
        'Otros': {'color': '1565C0'}
    }
    
    cols_order = [
        'Fecha', 'Exchange', 'Tipo de Operación', 'Moneda', 
        'Monto Compra (Cripto)', 'Monto Venta (Cripto)', 
        'Cotización Compra', 'Cotización Venta', 'Monto ARS', 'Comentarios'
    ]
    
    df_all = pd.DataFrame(transactions)
    for c in cols_order:
        if c not in df_all.columns:
            df_all[c] = ''
            
    if 'Exchange' not in df_all.columns or df_all.empty:
        return None

    # Exclude null/ghost transactions with 0.0 in all amount fields
    m_c = pd.to_numeric(df_all.get('Monto Compra (Cripto)', df_all.get('monto_compra_cripto', 0.0)), errors='coerce').fillna(0.0)
    m_v = pd.to_numeric(df_all.get('Monto Venta (Cripto)', df_all.get('monto_venta_cripto', 0.0)), errors='coerce').fillna(0.0)
    m_a = pd.to_numeric(df_all.get('Monto ARS', df_all.get('monto_ars', 0.0)), errors='coerce').fillna(0.0)
    df_all = df_all[(m_c.abs() > 1e-9) | (m_v.abs() > 1e-9) | (m_a.abs() > 1e-9)].copy()
        
    grouped_exchange = df_all.groupby('Exchange')
    current_row = 1
    
    for exchange_name, ex_group in grouped_exchange:
        ex_style = STYLES.get(str(exchange_name).split(' ')[0], STYLES.get(exchange_name, STYLES['Generic']))
        if 'Binance' in str(exchange_name): ex_style = STYLES['Binance Spot']
        if 'Ripio' in str(exchange_name): ex_style = STYLES['Ripio Trade']
        if 'Bitso' in str(exchange_name): ex_style = STYLES['Bitso']
        if 'Fiwind' in str(exchange_name): ex_style = STYLES['Fiwind']
        if 'OKX' in str(exchange_name): ex_style = STYLES['OKX']
        if 'Bybit' in str(exchange_name): ex_style = STYLES['Bybit']
        if 'Bitget' in str(exchange_name): ex_style = STYLES['Bitget']
        
        ws.cell(row=current_row, column=1, value=f"REPORTE: {exchange_name.upper()}")
        ws.cell(row=current_row, column=1).font = Font(bold=True, size=16, color=ex_style['font'])
        ws.cell(row=current_row, column=1).fill = PatternFill(start_color=ex_style['fill'], end_color=ex_style['fill'], fill_type="solid")
        current_row += 2
        
        def categorize_tipo(t):
            t_upper = str(t).upper()
            if 'COMPRA' in t_upper: return 'Compra'
            if 'VENTA' in t_upper: return 'Venta'
            return 'Otros'
            
        ex_group = ex_group.copy()
        ex_group['TipoCategoria'] = ex_group['Tipo de Operación'].apply(categorize_tipo)
        
        for cat in ['Compra', 'Venta', 'Otros']:
            cat_group = ex_group[ex_group['TipoCategoria'] == cat]
            if cat_group.empty:
                continue
            
            cat_group = cat_group.sort_values(by='Fecha')
            cat_color = TYPE_STYLES.get(cat, TYPE_STYLES['Otros'])['color']
            cat_title = "OTROS" if cat == "Otros" else f"{cat.upper()}S"
            ws.cell(row=current_row, column=1, value=cat_title)
            ws.cell(row=current_row, column=1).font = Font(bold=True, size=12, color=cat_color)
            current_row += 1
            
            for col_idx, col_name in enumerate(cols_order, 1):
                cell = ws.cell(row=current_row, column=col_idx, value=col_name)
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill(start_color='333333', end_color='333333', fill_type="solid")
                cell.alignment = Alignment(horizontal='center')
            current_row += 1
            
            start_data_row = current_row
            for _, row_data in cat_group.iterrows():
                for col_idx, col_name in enumerate(cols_order, 1):
                    val = row_data.get(col_name, '')
                    cell = ws.cell(row=current_row, column=col_idx, value=val)
                    cell.alignment = Alignment(horizontal='center')
                    if isinstance(val, (int, float)):
                        if 'Monto' in col_name or 'Cotización' in col_name:
                            cell.number_format = '#,##0.00'
                current_row += 1
                
            end_data_row = current_row - 1
            
            ws.cell(row=current_row, column=8, value="TOTAL M. ARS:")
            ws.cell(row=current_row, column=8).font = Font(bold=True)
            ws.cell(row=current_row, column=8).alignment = Alignment(horizontal='right')
            
            sum_cell = ws.cell(row=current_row, column=9)
            sum_cell.value = f"=SUBTOTAL(109, I{start_data_row}:I{end_data_row})"
            sum_cell.font = Font(bold=True, color=cat_color)
            sum_cell.number_format = '#,##0.00'
            sum_cell.alignment = Alignment(horizontal='center')
            
            border = Border(top=Side(style='thin', color='000000'), bottom=Side(style='double', color='000000'))
            sum_cell.border = border
            current_row += 3
            
        current_row += 5

    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if str(cell.value).startswith('='):
                    continue
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except Exception:
                pass
        ws.column_dimensions[column].width = min(max_length + 2, 50)

    wb.save(output)
    output.seek(0)
    return output

def generate_excel_bytes(transactions):
    """Backwards-compatible wrapper returning BytesIO of the master Excel."""
    return generate_master_excel(transactions)

def _parse_sheet_rows(rows, default_exchange='Otros'):
    """Helper state machine to parse a table of rows into transactions."""
    records = []
    current_exchange = default_exchange
    col_map = None

    for idx, row in enumerate(rows, start=1):
        try:
            if not row or not any(c is not None and str(c).strip() != '' for c in row):
                continue

            row_strs = [str(c).strip() if c is not None else '' for c in row]
            first_cell = row_strs[0] if row_strs else ''

            banner_match = re.search(r'REPORTE:\s*(.+)', first_cell, re.IGNORECASE)
            if not banner_match:
                for cell in row_strs:
                    bm = re.search(r'REPORTE:\s*(.+)', cell, re.IGNORECASE)
                    if bm:
                        banner_match = bm
                        break

            if banner_match:
                current_exchange = banner_match.group(1).strip()
                continue

            normalized_row = [_normalize_text(c) for c in row_strs]
            if any(h in normalized_row for h in ['fecha', 'date', 'timestamp', 'fecha/hora', 'datetime']):
                col_map = {}
                for col_i, norm_val in enumerate(normalized_row):
                    if not norm_val:
                        continue
                    if norm_val in ['fecha', 'date', 'timestamp', 'fecha/hora', 'datetime', 'created_at', 'hora'] and 'date' not in col_map:
                        col_map['date'] = col_i
                    elif norm_val in ['exchange', 'plataforma', 'entidad', 'broker', 'origen', 'cuenta', 'exchange/plataforma', 'plataforma/exchange'] and 'exchange' not in col_map:
                        col_map['exchange'] = col_i
                    elif norm_val in ['tipo', 'tipo de operacion', 'tipo de operación', 'type', 'operacion', 'operación', 'action', 'side'] and 'type' not in col_map:
                        col_map['type'] = col_i
                    elif norm_val in ['moneda', 'asset', 'cripto', 'ticker', 'coin', 'symbol', 'currency'] and 'asset' not in col_map:
                        col_map['asset'] = col_i
                    elif norm_val in ['monto compra (cripto)', 'monto compra', 'cantidad compra', 'monto_compra_cripto', 'buy_amount', 'compra'] and 'buy' not in col_map:
                        col_map['buy'] = col_i
                    elif norm_val in ['monto venta (cripto)', 'monto venta', 'cantidad venta', 'monto_venta_cripto', 'sell_amount', 'venta'] and 'sell' not in col_map:
                        col_map['sell'] = col_i
                    elif norm_val in ['cantidad', 'monto', 'amount', 'quantity', 'volume'] and 'amount' not in col_map:
                        col_map['amount'] = col_i
                    elif norm_val in ['monto ars', 'monto_ars', 'total ars', 'monto usd', 'total', 'total fiat', 'monto fiat', 'fiat', 'ars'] and 'fiat' not in col_map:
                        col_map['fiat'] = col_i
                    elif norm_val in ['cotización compra', 'cotizacion compra', 'cotización_compra', 'cotizacion_compra', 'precio compra'] and 'cot_buy' not in col_map:
                        col_map['cot_buy'] = col_i
                    elif norm_val in ['cotización venta', 'cotizacion venta', 'cotización_venta', 'cotizacion_venta', 'precio venta'] and 'cot_sell' not in col_map:
                        col_map['cot_sell'] = col_i
                    elif norm_val in ['precio', 'price', 'cotizacion', 'cotización', 'rate'] and 'price' not in col_map:
                        col_map['price'] = col_i
                    elif norm_val in ['comentarios', 'comentario', 'notes', 'memo', 'detalle', 'observaciones', 'referencia'] and 'notes' not in col_map:
                        col_map['notes'] = col_i
                    elif norm_val in ['id', 'order id', 'order_id', 'order no', 'order no.', 'txid', 'tx_hash', 'hash', 'identificador', 'ref'] and 'id' not in col_map:
                        col_map['id'] = col_i
                continue

            if not col_map:
                continue

            if len([c for c in row_strs if c]) == 1 and row_strs[0].upper() in ['COMPRAS', 'VENTAS', 'DEPOSITOS', 'RETIROS', 'TRANSACCIONES']:
                continue

            date_idx = col_map.get('date')
            if date_idx is None or date_idx >= len(row):
                continue

            raw_date = row[date_idx]
            if raw_date is None or str(raw_date).strip() == '' or str(raw_date).strip().lower() in ['fecha', 'date', 'nan', 'none']:
                continue

            if any('total' in str(c).lower() for c in row_strs[:4] if c):
                continue

            dt = parse_date(raw_date, current_exchange)

            ex_idx = col_map.get('exchange')
            raw_ex_val = str(row[ex_idx]).strip() if ex_idx is not None and ex_idx < len(row) and row[ex_idx] is not None else ''
            if raw_ex_val and raw_ex_val.lower() not in ('', 'nan', 'none', 'null'):
                row_ex = raw_ex_val
            else:
                row_ex = current_exchange if current_exchange and current_exchange.lower() not in ('', 'nan', 'none', 'null') else 'Otros'

            notes_idx = col_map.get('notes')
            raw_notes = str(row[notes_idx]).strip() if notes_idx is not None and notes_idx < len(row) and row[notes_idx] is not None and str(row[notes_idx]).strip() not in ['', 'nan', 'None'] else ''

            type_idx = col_map.get('type')
            raw_tipo_val = str(row[type_idx]).strip() if type_idx is not None and type_idx < len(row) and row[type_idx] is not None and str(row[type_idx]).strip() not in ['', 'nan', 'None', 'null'] else ''

            asset_idx = col_map.get('asset')
            raw_moneda = str(row[asset_idx]).strip().upper() if asset_idx is not None and asset_idx < len(row) and row[asset_idx] is not None and str(row[asset_idx]).strip() not in ['', 'nan', 'None', 'null'] else ''
            if not raw_moneda and raw_notes:
                for coin in ['USDT', 'USDC', 'BTC', 'ETH', 'DAI', 'SOL', 'BNB', 'XRP', 'ADA', 'DOGE', 'DOT', 'MATIC', 'AVAX', 'TRX', 'BUSD']:
                    if re.search(r'\b' + coin + r'\b', raw_notes.upper()):
                        raw_moneda = coin
            if not raw_moneda:
                import logging
                logging.getLogger(__name__).warning(f"Fila {idx} en multi-exchange omitida: moneda cripto vacía o indeterminada.")
                continue

            buy_idx = col_map.get('buy')
            sell_idx = col_map.get('sell')
            amount_idx = col_map.get('amount')

            m_compra = clean_decimal(row[buy_idx]) if buy_idx is not None and buy_idx < len(row) else 0.0
            m_venta = clean_decimal(row[sell_idx]) if sell_idx is not None and sell_idx < len(row) else 0.0

            if m_compra == 0.0 and m_venta == 0.0 and amount_idx is not None and amount_idx < len(row):
                amt = clean_decimal(row[amount_idx])
                if raw_tipo_val and any(kw in raw_tipo_val.upper() for kw in ['VENTA', 'RETIRO', 'SELL', 'OUT', 'EGRESO', 'ENVIO', 'ENVÍO', 'SALIDA', 'EXTRACCION', 'EXTRACCIÓN', 'COMISION', 'COMISIÓN']):
                    m_venta = amt
                else:
                    m_compra = amt

            if m_compra == 0.0 and m_venta == 0.0:
                continue

            if raw_tipo_val:
                if 'COMPRA' in raw_tipo_val.upper() or 'BUY' in raw_tipo_val.upper():
                    tipo_op = 'Compra'
                elif 'VENTA' in raw_tipo_val.upper() or 'SELL' in raw_tipo_val.upper():
                    tipo_op = 'Venta'
                else:
                    tipo_op = raw_tipo_val.title() if (raw_tipo_val.isupper() or raw_tipo_val.islower()) else raw_tipo_val
            else:
                if m_venta > 0 and m_compra == 0:
                    tipo_op = 'Venta'
                else:
                    tipo_op = 'Compra'

            cot_buy_idx = col_map.get('cot_buy')
            cot_sell_idx = col_map.get('cot_sell')
            price_idx = col_map.get('price')

            c_compra = clean_decimal(row[cot_buy_idx]) if cot_buy_idx is not None and cot_buy_idx < len(row) else (clean_decimal(row[price_idx]) if price_idx is not None and price_idx < len(row) and m_compra > 0 else 0.0)
            c_venta = clean_decimal(row[cot_sell_idx]) if cot_sell_idx is not None and cot_sell_idx < len(row) else (clean_decimal(row[price_idx]) if price_idx is not None and price_idx < len(row) and m_venta > 0 else 0.0)

            fiat_idx = col_map.get('fiat')
            m_ars = clean_decimal(row[fiat_idx]) if fiat_idx is not None and fiat_idx < len(row) else 0.0

            if m_ars == 0.0:
                coin_norm = raw_moneda.split('/')[0].split('-')[0].split('_')[0].strip().upper()
                if m_compra > 0 and c_compra > 0:
                    usd_ars_rate = _get_usd_ars_rate_for_date(dt)
                    is_usd_scale = usd_ars_rate > 100.0 and (
                        (coin_norm in _STABLES and c_compra < 15.0)
                        or (coin_norm in _REF_USD and c_compra < _REF_USD[coin_norm] * 5.0)
                    )
                    if is_usd_scale:
                        c_compra = c_compra * usd_ars_rate
                    m_ars = m_compra * c_compra
                elif m_venta > 0 and c_venta > 0:
                    usd_ars_rate = _get_usd_ars_rate_for_date(dt)
                    is_usd_scale = usd_ars_rate > 100.0 and (
                        (coin_norm in _STABLES and c_venta < 15.0)
                        or (coin_norm in _REF_USD and c_venta < _REF_USD[coin_norm] * 5.0)
                    )
                    if is_usd_scale:
                        c_venta = c_venta * usd_ars_rate
                    m_ars = m_venta * c_venta
                elif '[INTERCAMBIO:' in raw_notes:
                    from models.transaction_models import get_quote_usd_value
                    usd_rate = _get_usd_ars_rate_for_date(dt)
                    qty = m_compra if m_compra > 0 else m_venta
                    m_ars = get_quote_usd_value(raw_moneda, qty, fecha=dt) * usd_rate
                    if m_compra > 0 and c_compra == 0:
                        c_compra = m_ars / m_compra if m_compra > 0 else 0.0
                    elif m_venta > 0 and c_venta == 0:
                        c_venta = m_ars / m_venta if m_venta > 0 else 0.0

            id_idx = col_map.get('id')
            raw_id = str(row[id_idx]).strip() if id_idx is not None and id_idx < len(row) and row[id_idx] is not None and str(row[id_idx]).strip() not in ['', 'nan', 'None', 'null'] else None
            if not raw_id and raw_notes:
                id_m = re.search(r'\bID:\s*([a-zA-Z0-9_\-]+)', raw_notes, re.IGNORECASE)
                if id_m:
                    raw_id = id_m.group(1).strip()

            tx_dict = create_transaction(
                dt, row_ex, tipo_op, raw_moneda,
                m_compra, m_venta, c_compra, c_venta, m_ars,
                comentario=raw_notes,
                unique_id=raw_id
            )
            if tx_dict:
                records.append(tx_dict)
        except Exception as row_err:
            print(f"Error procesando fila {idx} en multi exchange excel: {row_err}")
            continue

    return records

def process_multi_exchange_excel(file_obj, filename):
    """
    Parses a multi-exchange manual Excel/CSV file using single-pass state machine.
    Supports multi-tab Excel files where each tab is an exchange.
    Returns (records, sample_records).
    """
    file_obj.seek(0)
    ext = filename.lower()
    records = []

    try:
        if ext.endswith('.xlsx'):
            wb = openpyxl.load_workbook(file_obj, data_only=True)
            for ws in wb.worksheets:
                rows = list(ws.iter_rows(values_only=True))
                if rows:
                    tab_name = ws.title.strip()
                    def_ex = 'Otros' if tab_name.lower().startswith(('sheet', 'hoja', 'tabla')) else tab_name
                    sheet_records = _parse_sheet_rows(rows, default_exchange=def_ex)
                    records.extend(sheet_records)
        elif ext.endswith('.xls'):
            try:
                xls_dict = pd.read_excel(file_obj, sheet_name=None, header=None)
                for tab_name, s_df in xls_dict.items():
                    if not s_df.empty:
                        rows = [tuple(r) for r in s_df.itertuples(index=False, name=None)]
                        def_ex = 'Otros' if str(tab_name).lower().startswith(('sheet', 'hoja', 'tabla')) else str(tab_name).strip()
                        sheet_records = _parse_sheet_rows(rows, default_exchange=def_ex)
                        records.extend(sheet_records)
            except Exception as e:
                print(f"Error reading .xls file with pandas: {e}")
        else:
            import csv
            for enc in ['utf-8-sig', 'utf-8', 'latin1', 'cp1252']:
                file_obj.seek(0)
                try:
                    content = file_obj.read()
                    if isinstance(content, bytes):
                        content = content.decode(enc)
                    dialect = None
                    try:
                        dialect = csv.Sniffer().sniff(content[:2048])
                    except Exception:
                        pass
                    sep = dialect.delimiter if dialect else ','
                    reader = csv.reader(io.StringIO(content), delimiter=sep)
                    rows = [r for r in reader]
                    if rows:
                        records = _parse_sheet_rows(rows, default_exchange='Otros')
                        break
                except Exception:
                    continue
    except Exception as e:
        print(f"Error reading multi-exchange file: {e}")
        return [], []

    sample = records[:5]
    return records, sample
